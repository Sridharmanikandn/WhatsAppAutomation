import requests
import pywhatkit
import time
from openpyxl import load_workbook
import os

def download_image(url, image_path):
    try:
        response = requests.get(url, stream=True)
        if response.status_code == 200:
            with open(image_path, "wb") as file:
                for chunk in response.iter_content(1024):
                    file.write(chunk)
            print(f"Image downloaded successfully to {image_path}!")
            return True
        else:
            print(f"Failed to download the image from {url}. Status code: {response.status_code}")
            return False
    except requests.exceptions.RequestException as e:
        print(f"A request error occurred while downloading the image from {url}: {e}")
        return False
    except Exception as e:
        print(f"An unexpected error occurred while downloading the image from {url}: {e}")
        return False

def send_image_with_caption(phone_number, image_path, caption):
    try:
        print(f"Initiating sending image {image_path} to {phone_number}...")
        pywhatkit.sendwhats_image(phone_number, image_path, caption=caption, wait_time=20, tab_close=True)
        print(f"Image {image_path} sent successfully to {phone_number}!")
        return True
    except Exception as e:
        print(f"An error occurred while sending the image to {phone_number}: {e}")
        return False

def clear_images(customer_name):
    try:
        before_path = f"before_{customer_name}.png"
        after_path = f"after_{customer_name}.png"
        if os.path.exists(before_path):
            os.remove(before_path)
            print(f"Removed {before_path}")
        if os.path.exists(after_path):
            os.remove(after_path)
            print(f"Removed {after_path}")
    except Exception as e:
        print(f"Error while clearing images: {e}")

filename = "car (1).xlsx"  
try:
    workbook = load_workbook(filename)
    sheet = workbook.active
except FileNotFoundError:
    print(f"Error: File '{filename}' not found.")
    exit()
except Exception as e:
    print(f"Error loading workbook: {e}")
    exit()

for row in sheet.iter_rows(min_row=2):
    try:
        phone_number = row[0].value
        customer_name = row[1].value
        before_url = row[2].value
        after_url = row[3].value

        if not phone_number or not before_url or not after_url:
            print(f"Skipping row with missing data: {row}")
            continue

      
        if isinstance(phone_number, float): 
            phone_number = int(phone_number)
        
        if not isinstance(phone_number, str):
            phone_number = str(phone_number)

        phone_number = phone_number.replace(" ", "").replace("-", "")

        if not phone_number.startswith("+"):
            if not phone_number.startswith("91"):
                phone_number = "+91" + phone_number

        print(f"Processing: Phone: {phone_number}, Name: {customer_name}, Before: {before_url}, After: {after_url}")

        download_successful_before = download_image(before_url, f"before_{customer_name}.png")
        download_successful_after = download_image(after_url, f"after_{customer_name}.png")

        if download_successful_before and download_successful_after:
            caption = f"Hi {customer_name}, here are your images!"
            send_successful_before = send_image_with_caption(phone_number, f"before_{customer_name}.png", caption)
            send_successful_after = send_image_with_caption(phone_number, f"after_{customer_name}.png", caption)

            if not send_successful_before or not send_successful_after:
                print(f"Error sending to {phone_number}. Check WhatsApp Web, number validity, and internet connection.")
            else:
                time.sleep(15)
        else:
            print(f"Download failed for {customer_name}. Check URLs and internet connection.")

        clear_images(customer_name)
    except Exception as e:
        print(f"An error occurred while processing a row: {e}")
        clear_images(customer_name)
        continue

print("Finished processing the Excel file.")